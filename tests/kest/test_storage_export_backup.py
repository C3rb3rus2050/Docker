"""Dokumentenablage, Excel-Export, Backup/Wiederherstellung, Konfiguration."""

from __future__ import annotations

import zipfile
from datetime import date

import pytest
from openpyxl import load_workbook

from broker_collector.security import UnsafePathError
from kest_helpers import APPLE, MSCI_WORLD_ACC, REIT, add_tx, fund, stock, ctx, depot, repo  # noqa: F401
from kest_tool.backup import BackupError, create_backup, restore_backup, verify_backup
from kest_tool.config import ConfigError, load_settings
from kest_tool.context import create_context
from kest_tool.db import Database
from kest_tool.export.excel import SHEETS, export_excel
from kest_tool.export.reports import export_csv, export_e1kv_helper, export_pdf_report
from kest_tool.models import FundReportIn, FundStatus, SecurityType, TxType
from kest_tool.tax.engine import compute


# --------------------------------------------------------------------------- #
# Dokumentenablage
# --------------------------------------------------------------------------- #
def test_folder_structure_created(ctx):
    base = ctx.store.base
    for sub in ("database/kest_tool.db", "config", "logs", "documents/2024/TradeRepublic",
                "documents/2025/ScalableCapital", "documents/2026/OeKB", "documents/2024/Other",
                "documents/2024/Unclassified", "imports/2024/original", "imports/2024/processed",
                "imports/2024/import_logs", "exports/2024/excel", "exports/2024/pdf", "exports/2024/csv",
                "exports/2024/e1kv", "reports/2024"):
        assert (base / sub).exists(), sub


def test_document_storage_and_retrieval(repo):
    data = b"%PDF-1.4 Beispiel A %%EOF"
    doc, new = repo.add_document(data, "Abrechnung.pdf", doc_type="Kaufabrechnung", source="TradeRepublic",
                                 tax_year=2024)
    assert new
    path = repo.document_path(doc)
    assert path.read_bytes() == data
    assert path.parent == repo.store.base / "documents" / "2024" / "TradeRepublic"
    # gleiche Datei erneut → keine Kopie
    doc2, new2 = repo.add_document(data, "anderer_name.pdf", source="TradeRepublic", tax_year=2024)
    assert not new2 and doc2["id"] == doc["id"]
    # gleicher Name, anderer Inhalt → kein Überschreiben
    doc3, new3 = repo.add_document(b"%PDF-1.4 Beispiel B %%EOF", "Abrechnung.pdf", source="TradeRepublic",
                                   tax_year=2024)
    assert new3 and repo.document_path(doc3).name == "Abrechnung (2).pdf"
    assert path.read_bytes() == data
    # Wiederauffinden
    assert [d["id"] for d in repo.documents(tax_year=2024, source="TradeRepublic", search="Abrechnung")] == \
        [doc3["id"], doc["id"]]
    assert repo.documents(doc_type="Kaufabrechnung")[0]["id"] == doc["id"]
    assert repo.find_document_by_hash(doc["sha256"])["id"] == doc["id"]


def test_document_recategorize_moves_without_changing_content(repo):
    data = b"%PDF-1.4 unklar %%EOF"
    doc, _ = repo.add_document(data, "scan.pdf")
    assert "/Unclassified/" in "/" + doc["stored_path"]
    repo.update_document(doc["id"], tax_year=2025, source="ScalableCapital", doc_type="Dividendenabrechnung")
    moved = repo.document(doc["id"])
    assert moved["stored_path"].startswith("documents/2025/ScalableCapital/")
    assert repo.document_path(moved).read_bytes() == data
    assert any(e["event"] == "geändert" for e in repo.document_events(doc["id"]))


def test_document_transaction_link(repo, depot):
    stock(repo)
    doc, _ = repo.add_document(b"%PDF-1.4 beleg %%EOF", "beleg.pdf", tax_year=2024)
    tid = add_tx(repo, depot, TxType.DIVIDEND, "2024-05-15", isin=APPLE, amount=10)
    repo.link_document(doc["id"], tid)
    assert [t["id"] for t in repo.transactions_for_document(doc["id"])] == [tid]
    assert repo.documents()[0]["tx_count"] == 1


def test_path_manipulation_rejected(repo):
    with pytest.raises(UnsafePathError):
        repo.store.absolute("../../etc/passwd")
    with pytest.raises(UnsafePathError):
        repo.store.absolute("documents/../../x")
    doc, _ = repo.add_document(b"data", "..\\..\\Windows\\CON.pdf", tax_year=2024)
    p = repo.document_path(doc)
    assert p.is_relative_to(repo.store.base) and p.name.startswith("_CON")


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def _scenario(repo, depot):
    stock(repo)
    stock(repo, isin=REIT, name="Realty Income", stype=SecurityType.REIT, verified=False)
    fund(repo, MSCI_WORLD_ACC, FundStatus.MELDEFONDS, name="MSCI World")
    add_tx(repo, depot, TxType.BUY, "2024-01-10", isin=APPLE, quantity=10, amount=1000, fees=1)
    add_tx(repo, depot, TxType.SELL, "2024-06-10", isin=APPLE, quantity=5, amount=400, fees=1)
    add_tx(repo, depot, TxType.DIVIDEND, "2024-05-15", isin=APPLE, amount=20, withholding_tax=3,
           withholding_country="US")
    add_tx(repo, depot, TxType.DIVIDEND, "2024-02-15", isin=REIT, amount=25, withholding_tax=3.75,
           withholding_country="US")
    add_tx(repo, depot, TxType.BUY, "2024-01-10", isin=MSCI_WORLD_ACC, quantity=10, amount=800)
    add_tx(repo, depot, TxType.INTEREST, "2024-12-31", amount=12.34)
    add_tx(repo, depot, TxType.DIVIDEND, "2023-05-15", isin=APPLE, amount=7)
    repo.add_fund_report(FundReportIn(isin=MSCI_WORLD_ACC, report_date=date(2024, 6, 28), age_per_unit=1.2345,
                                      ak_correction_per_unit=1.1, source="manuell"))


def test_excel_contains_all_sheets(repo, depot, ctx):
    _scenario(repo, depot)
    res = compute(repo, 2024)
    path = export_excel(repo, ctx.store, res, today=date(2025, 1, 15))
    assert path.name == "KESt_Auswertung_2024_2025-01-15.xlsx"
    assert path.parent == ctx.store.base / "exports" / "2024" / "excel"
    wb = load_workbook(path)
    assert wb.sheetnames == SHEETS
    for name in SHEETS:
        ws = wb[name]
        assert ws.freeze_panes is not None, name
        assert ws.auto_filter.ref, name
    assert repo.exports(2024)[0]["kind"] == "excel"
    # zweiter Export am selben Tag überschreibt nicht
    path2 = export_excel(repo, ctx.store, res, today=date(2025, 1, 15))
    assert path2 != path and path.exists()


def _sheet_table(ws, header_row: int = 4):
    headers = [c.value for c in ws[header_row]]
    rows = []
    for r in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if r[0] is None or (isinstance(r[0], str) and r[0].startswith("Summe (SUBTOTAL")):
            break
        rows.append(dict(zip(headers, r)))
    return rows


def test_excel_sums_match_database(repo, depot, ctx):
    _scenario(repo, depot)
    res = compute(repo, 2024)
    wb = load_workbook(export_excel(repo, ctx.store, res))
    tx_rows = _sheet_table(wb["Alle Transaktionen"])
    db_tx = repo.transactions()
    assert len(tx_rows) == len(db_tx)
    assert sum(r["Betrag EUR"] or 0 for r in tx_rows) == pytest.approx(sum(t["amount_eur"] or 0 for t in db_tx))
    assert sum(r["Gebühren EUR"] or 0 for r in tx_rows) == pytest.approx(sum(t["fees_eur"] or 0 for t in db_tx))
    assert sum(r["Quellensteuer EUR"] or 0 for r in tx_rows) == pytest.approx(
        sum(t["withholding_tax_eur"] or 0 for t in db_tx))
    # Summenzeile als nachvollziehbare Formel
    ws = wb["Alle Transaktionen"]
    total_row = 4 + len(tx_rows) + 1
    assert str(ws.cell(row=total_row, column=16).value).startswith("=SUBTOTAL(9,P5:P")
    # Übersicht = Steuerberechnung
    summary = {r["Kennzahl"]: r["Gesamt (inkl. vorläufig)"] for r in _sheet_table(wb["Steuerübersicht"], 5)}
    for k in ("Dividenden", "Realisierte Verluste", "Ausschüttungsgleiche Erträge", "Steuer 27,5 %",
              "Steuer 25 %", "Bemessungsgrundlage 27,5 %"):
        assert summary[k] == pytest.approx(res.summary[k]), k
    calc = _sheet_table(wb["Steuerberechnung"])
    assert len(calc) == len(res.lines)
    assert sum(r["Bemessungsgrundlage"] or 0 for r in calc if r["Steuergruppe"] == "27,5 %") == \
        pytest.approx(res.summary["Bemessungsgrundlage 27,5 %"])
    div = _sheet_table(wb["Dividenden und Ausschüttungen"])
    assert sum(r["Bruttoertrag EUR"] for r in div if r["Steuerliche Kategorie"] == "Dividende") == \
        pytest.approx(res.summary["Dividenden"])
    # REIT ohne Klassifizierung: als ungeklärt sichtbar, nicht als Dividende
    reit = _sheet_table(wb["REITs und BDCs"])
    assert reit[0]["Steuerlicher Ansatz"] == "nicht berechnet"
    assert "UNGEPRÜFT" in reit[0]["Steuerliche Klassifizierung"]
    e1 = _sheet_table(wb["E1kv - FinanzOnline"])
    assert {r["Kennzahl"]: r["Betrag EUR"] for r in e1}["863"] == pytest.approx(20)
    prot = _sheet_table(wb["Prüfprotokoll"])
    assert any("REIT" in (r["Meldung"] or "") for r in prot)


def test_other_exports(repo, depot, ctx):
    _scenario(repo, depot)
    res = compute(repo, 2024)
    csvs = export_csv(repo, ctx.store, res)
    assert len(csvs) == 4 and all(p.exists() for p in csvs)
    txt = export_e1kv_helper(repo, ctx.store, res)
    content = txt.read_text(encoding="utf-8")
    assert "KEINE OFFIZIELLE EINREICHUNG" in content and "863" in content
    pdf = export_pdf_report(repo, ctx.store, res)
    assert pdf.read_bytes().startswith(b"%PDF")
    assert list((ctx.store.base / "reports" / "2024").glob("*.pdf"))


# --------------------------------------------------------------------------- #
# Backup
# --------------------------------------------------------------------------- #
def test_backup_and_restore(repo, depot, ctx, tmp_path):
    _scenario(repo, depot)
    repo.add_document(b"%PDF-1.4 original %%EOF", "orig.pdf", tax_year=2024, source="TradeRepublic")
    zip_path = create_backup(ctx.db, ctx.store)
    manifest = verify_backup(zip_path)
    assert "database/kest_tool.db" in manifest["files"]
    target = tmp_path / "Wiederhergestellt"
    info = restore_backup(zip_path, target)
    assert info["files"] == len(manifest["files"])
    db2 = Database(target / "database" / "kest_tool.db")
    assert db2.one("SELECT COUNT(*) n FROM transactions")["n"] == len(repo.transactions())
    doc = db2.one("SELECT * FROM documents WHERE original_filename='orig.pdf'")
    assert (target / doc["stored_path"]).read_bytes() == b"%PDF-1.4 original %%EOF"
    db2.close()
    # nicht leerer Zielordner nur mit Verschieben
    with pytest.raises(BackupError):
        restore_backup(zip_path, target)
    info2 = restore_backup(zip_path, target, move_existing=True)
    assert info2["moved_existing_to"] and (tmp_path / info2["moved_existing_to"]).exists()


def test_backup_tamper_and_zip_slip(ctx, tmp_path):
    zip_path = create_backup(ctx.db, ctx.store)
    bad = tmp_path / "bad.zip"
    with zipfile.ZipFile(zip_path) as src, zipfile.ZipFile(bad, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "database/kest_tool.db":
                data = data[:-10] + b"0123456789"
            dst.writestr(item, data)
    with pytest.raises(BackupError, match="Prüfsumme"):
        verify_backup(bad)
    slip = tmp_path / "slip.zip"
    with zipfile.ZipFile(slip, "w") as z:
        z.writestr("manifest.json", '{"created_at": "x", "files": {"../evil.txt": "00"}}')
        z.writestr("../evil.txt", "boom")
    with pytest.raises(BackupError, match="Unzulässiger Pfad"):
        verify_backup(slip)
    assert not (tmp_path.parent / "evil.txt").exists()


# --------------------------------------------------------------------------- #
# Konfiguration
# --------------------------------------------------------------------------- #
def test_config_rejects_secrets(tmp_path):
    f = tmp_path / "c.yaml"
    f.write_text("broker:\n  password: geheim\n", encoding="utf-8")
    with pytest.raises(ConfigError):
        load_settings(default_file=f, local_file=tmp_path / "none.yaml")


def test_config_base_dir_and_local_override(tmp_path, monkeypatch):
    monkeypatch.delenv("KEST_TOOL_HOME", raising=False)
    monkeypatch.setenv("USERPROFILE", str(tmp_path / "home"))
    s = load_settings(local_file=tmp_path / "local.yaml")
    assert s.base_dir == tmp_path / "home" / "Documents" / "KESt-Tool"
    s.save_local({"base_dir": str(tmp_path / "Ablage")})
    s2 = load_settings(local_file=tmp_path / "local.yaml")
    assert s2.base_dir == tmp_path / "Ablage"
    ctx = create_context(s2)
    assert (tmp_path / "Ablage" / "config" / "config.yaml").exists()
    ctx.close()
