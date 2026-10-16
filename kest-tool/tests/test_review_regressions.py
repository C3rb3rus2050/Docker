"""Regressionen zu Quellenstand, Datenherkunft, Sortierung und Export-Sicherheit."""
from datetime import date
import csv
import math

import pytest
from openpyxl import Workbook, load_workbook

from kest_helpers import APPLE, MSCI_WORLD_ACC, NMF_FUND, add_tx, stock, fund, ctx, repo, depot  # noqa: F401
from kest_tool.models import FundReportIn, FundStatus, SecurityIn, SecurityType, TaxClass, TransactionIn, TxType
from kest_tool.tax.engine import CAT_GAIN, CAT_NMF, MANUAL, PROVISIONAL, compute
from kest_tool.tax.rules import load_rules
from kest_tool.export.excel import Col, SheetWriter
from kest_tool.export.reports import _write_csv


def gain(result):
    return next(line for line in result.lines if line.category == CAT_GAIN)


@pytest.mark.parametrize('sell_first', [True, False])
def test_same_day_missing_order_never_guesses(repo, depot, sell_first):
    stock(repo)
    add_tx(repo, depot, TxType.BUY, '2024-01-01', isin=APPLE, quantity=10, amount=1000)
    txs = [(TxType.SELL, 5, 1000), (TxType.BUY, 10, 3000)]
    for typ, qty, amount in txs if sell_first else reversed(txs):
        add_tx(repo, depot, typ, '2024-02-01', isin=APPLE, quantity=qty, amount=amount)
    result = compute(repo, 2024)
    assert gain(result).amount is None and gain(result).status == MANUAL
    assert any('Reihenfolge fehlt' in issue.message for issue in result.issues)


@pytest.mark.parametrize('sell_first', [True, False])
def test_explicit_execution_order_is_independent_of_import(repo, depot, sell_first):
    stock(repo)
    add_tx(repo, depot, TxType.BUY, '2024-01-01', isin=APPLE, quantity=10, amount=1000)
    txs = [(TxType.SELL, 5, 1000, 2), (TxType.BUY, 10, 3000, 1)]
    for typ, qty, amount, order in txs if sell_first else reversed(txs):
        add_tx(repo, depot, typ, '2024-02-01', isin=APPLE, quantity=qty, amount=amount,
               execution_order=order)
    assert gain(compute(repo, 2024)).amount == pytest.approx(0)


def test_unreviewed_fund_costs_are_visible_on_later_sale(repo, depot):
    fund(repo, MSCI_WORLD_ACC, FundStatus.MELDEFONDS)
    add_tx(repo, depot, TxType.BUY, '2024-01-01', isin=MSCI_WORLD_ACC, quantity=10, amount=1000)
    report = repo.add_fund_report(FundReportIn(isin=MSCI_WORLD_ACC, report_date=date(2024, 6, 1),
                                 age_per_unit=1, ak_correction_per_unit=1))
    add_tx(repo, depot, TxType.SELL, '2025-06-01', isin=MSCI_WORLD_ACC, quantity=5, amount=600)
    result = compute(repo, 2025)
    assert gain(result).amount == pytest.approx(95)
    assert gain(result).status == PROVISIONAL
    assert any('ungeprüfter Fondsmeldung' in note for note in gain(result).notes)
    assert result.positions[0]['AK vollständig'] == 'nein'
    repo.set_fund_report_review(report['id'], 'geprueft')
    result = compute(repo, 2025)
    assert not any('ungeprüfter Fondsmeldung' in note for note in gain(result).notes)


def test_missing_historical_fund_costs_block_sale(repo, depot):
    fund(repo, MSCI_WORLD_ACC, FundStatus.MELDEFONDS)
    add_tx(repo, depot, TxType.BUY, '2023-01-01', isin=MSCI_WORLD_ACC, quantity=10, amount=1000)
    add_tx(repo, depot, TxType.SELL, '2025-06-01', isin=MSCI_WORLD_ACC, quantity=5, amount=600)
    line = gain(compute(repo, 2025))
    assert line.amount is None and line.status == MANUAL
    assert any('Historische Fondsdaten' in note for note in line.notes)


def test_past_corporate_action_blocks_later_sale(repo, depot):
    stock(repo)
    add_tx(repo, depot, TxType.BUY, '2023-01-01', isin=APPLE, quantity=10, amount=1000)
    add_tx(repo, depot, TxType.CORPORATE_ACTION, '2023-06-01', isin=APPLE, notes='Abspaltung')
    add_tx(repo, depot, TxType.SELL, '2024-06-01', isin=APPLE, quantity=5, amount=600)
    assert gain(compute(repo, 2024)).amount is None


def test_legacy_acquisition_is_not_silently_taxed_as_new_asset(repo, depot):
    stock(repo)
    add_tx(repo, depot, TxType.BUY, '2010-01-01', isin=APPLE, quantity=10, amount=1000)
    add_tx(repo, depot, TxType.SELL, '2024-06-01', isin=APPLE, quantity=5, amount=600)
    line = gain(compute(repo, 2024))
    assert line.amount is None and any('Alt-/Übergangsbestand' in n for n in line.notes)


@pytest.mark.parametrize('year', [2024, 2025, 2026])
def test_current_nonreporting_fund_rule_and_cost_rollforward(repo, depot, year):
    fund(repo, NMF_FUND, FundStatus.NICHTMELDEFONDS)
    add_tx(repo, depot, TxType.BUY, f'{year}-01-02', isin=NMF_FUND, quantity=10, amount=1000)
    repo.upsert_nmf_price(NMF_FUND, year, 100, f'{year}-01-02', 130, f'{year}-12-31', source='NAV')
    result = compute(repo, year)
    assert next(l for l in result.lines if l.category == CAT_NMF).amount == pytest.approx(270)
    if year < 2026:
        add_tx(repo, depot, TxType.SELL, f'{year+1}-06-01', isin=NMF_FUND, quantity=5, amount=700)
        assert gain(compute(repo, year+1)).amount == pytest.approx(65)


@pytest.mark.parametrize('country', ['US', 'DE', 'CH', 'NL', 'CA'])
def test_verified_ordinary_dividend_caps(repo, depot, country):
    stock(repo, country=country)
    add_tx(repo, depot, TxType.DIVIDEND, '2025-06-01', isin=APPLE, amount=100,
           withholding_tax=35, withholding_country=country)
    line = compute(repo, 2025).lines[0]
    assert line.creditable_tax_max == 15 and line.non_creditable_tax == 20


def test_interest_does_not_use_dividend_treaty_rate(repo, depot):
    add_tx(repo, depot, TxType.INTEREST, '2025-06-01', amount=100,
           withholding_tax=30, withholding_country='DE')
    line = compute(repo, 2025).lines[0]
    assert line.creditable_tax_max == 0 and line.unclear_tax == 30


def test_e1kv_2025_uses_correct_withholding_columns(repo, depot):
    stock(repo)
    domestic = repo.add_depot(repo.depots()[0]['broker_id'], 'AT')
    repo.set_depot_tax_settings(domestic, 2025, 'ja', 'ja', source='Test')
    for dep in (depot, domestic):
        add_tx(repo, dep, TxType.DIVIDEND, '2025-06-01', isin=APPLE, amount=100,
               withholding_tax=15, withholding_country='US')
    rows = {r['Kennzahl']: r for r in compute(repo, 2025).e1kv}
    assert rows['984']['Betrag EUR'] == 15
    assert rows['998']['Betrag EUR'] == 15
    assert rows['863']['Kennzahl-Status'] == 'offiziell belegt (Auszug)'
    assert load_rules().e1kv(2026).get('unverified_reference_year') == 2025


def test_excel_imported_text_is_literal_but_subtotal_is_formula(tmp_path):
    wb = Workbook()
    writer = SheetWriter(wb.active, 'Test')
    writer.table([Col('Name', 'name'), Col('Betrag', 'amount', 'eur', total=True)],
                 [{'name': '=HYPERLINK("https://example.org","test")', 'amount': 100}])
    path = tmp_path / 'export.xlsx'
    wb.save(path)
    result = load_workbook(path)
    assert result.active['A5'].data_type == 's'
    assert result.active['A5'].value.startswith('=HYPERLINK')
    assert result.active['B6'].data_type == 'f'


@pytest.mark.parametrize('value', ['=1+1', '+cmd', '-cmd', '@SUM(A1)', '  =1', '\t=1'])
def test_csv_text_formula_escaped_without_changing_numeric_values(tmp_path, value):
    path = tmp_path / 'export.csv'
    _write_csv(path, [{'text': value, 'amount': -12.5}])
    with path.open(encoding='utf-8-sig', newline='') as fh:
        row = next(csv.DictReader(fh, delimiter=';'))
    assert row['text'] == "'" + value
    assert row['amount'] == '-12,5'


@pytest.mark.parametrize('amount', [math.nan, math.inf, -math.inf])
def test_transactions_reject_nonfinite_amounts(depot, amount):
    with pytest.raises(ValueError):
        TransactionIn(depot_id=depot, tx_type=TxType.DIVIDEND, trade_date=date(2025, 1, 1), amount=amount)


@pytest.mark.parametrize('proceeds,kz,expected', [(1200, '995', 200), (800, '896', -200)])
def test_securitized_derivative_has_own_e1kv_fields(repo, depot, proceeds, kz, expected):
    repo.upsert_security(SecurityIn(isin=APPLE, security_type=SecurityType.CERTIFICATE,
                         tax_class=TaxClass.DERIVAT_VERBRIEFT, tax_class_verified=True,
                         tax_class_source='Test: Emittentenprospekt bestätigt Verbriefung und öffentliches Angebot'))
    add_tx(repo, depot, TxType.BUY, '2025-01-01', isin=APPLE, quantity=10, amount=1000)
    add_tx(repo, depot, TxType.SELL, '2025-06-01', isin=APPLE, quantity=10, amount=proceeds)
    result = compute(repo, 2025)
    assert result.lines[0].amount == expected
    rows = {r['Kennzahl']: r for r in result.e1kv}
    assert rows[kz]['Betrag EUR'] == expected
    assert rows['994']['Betrag EUR'] == 0
    assert rows['892']['Betrag EUR'] == 0


def test_generic_derivative_remains_manual(repo, depot):
    repo.upsert_security(SecurityIn(isin=APPLE, security_type=SecurityType.CERTIFICATE,
                         tax_class=TaxClass.DERIVAT))
    add_tx(repo, depot, TxType.BUY, '2025-01-01', isin=APPLE, quantity=10, amount=1000)
    add_tx(repo, depot, TxType.SELL, '2025-06-01', isin=APPLE, quantity=10, amount=1200)
    assert gain(compute(repo, 2025)).amount is None


def test_equity_acquired_in_2011_is_new_asset(repo, depot):
    stock(repo)
    add_tx(repo, depot, TxType.BUY, '2011-06-01', isin=APPLE, quantity=10, amount=1000)
    add_tx(repo, depot, TxType.SELL, '2025-06-01', isin=APPLE, quantity=10, amount=1200)
    assert gain(compute(repo, 2025)).amount == 200


def test_nonreporting_fund_year_end_position_includes_cost_increase(repo, depot):
    fund(repo, NMF_FUND, FundStatus.NICHTMELDEFONDS)
    add_tx(repo, depot, TxType.BUY, '2025-01-02', isin=NMF_FUND, quantity=10, amount=1000)
    repo.upsert_nmf_price(NMF_FUND, 2025, 100, '2025-01-02', 130, '2025-12-31', source='NAV')
    assert compute(repo, 2025).positions[0]['Steuerliche AK gesamt EUR'] == 1270


def test_unknown_cost_capital_return_does_not_create_false_gain(repo, depot):
    stock(repo)
    add_tx(repo, depot, TxType.BUY, '2025-01-02', isin=APPLE, quantity=10, amount=1000,
           currency='USD')
    add_tx(repo, depot, TxType.CAPITAL_REPAYMENT, '2025-06-01', isin=APPLE, amount=100,
           income_category_verified=True)
    result = compute(repo, 2025)
    assert all(l.amount is None for l in result.lines)


def test_historical_nonreporting_fx_missing_blocks_sale(repo, depot):
    fund(repo, NMF_FUND, FundStatus.NICHTMELDEFONDS)
    add_tx(repo, depot, TxType.BUY, '2024-01-02', isin=NMF_FUND, quantity=10, amount=1000)
    repo.upsert_nmf_price(NMF_FUND, 2024, 100, '2024-01-02', 130, '2024-12-31', currency='USD', source='NAV')
    add_tx(repo, depot, TxType.SELL, '2025-06-01', isin=NMF_FUND, quantity=10, amount=1500)
    assert gain(compute(repo, 2025)).amount is None


def test_closed_legacy_position_does_not_taint_new_acquisition(repo, depot):
    stock(repo)
    add_tx(repo, depot, TxType.BUY, '2010-01-01', isin=APPLE, quantity=10, amount=1000)
    add_tx(repo, depot, TxType.SELL, '2023-01-01', isin=APPLE, quantity=10, amount=1200)
    add_tx(repo, depot, TxType.BUY, '2024-01-01', isin=APPLE, quantity=10, amount=1500)
    add_tx(repo, depot, TxType.SELL, '2025-01-01', isin=APPLE, quantity=10, amount=1700)
    assert gain(compute(repo, 2025)).amount == 200


def test_execution_order_migration_preserves_existing_data(ctx, depot):
    from kest_tool.db import Database
    tid = add_tx(ctx.repo, depot, TxType.BUY, '2025-01-02', isin=APPLE, quantity=1, amount=100)
    ctx.repo.db.execute('ALTER TABLE transactions DROP COLUMN execution_order')
    reopened = Database(ctx.repo.db.path)
    try:
        row = reopened.one('SELECT * FROM transactions WHERE id=?', (tid,))
        assert row['amount'] == 100 and row['execution_order'] is None
    finally:
        reopened.close()
