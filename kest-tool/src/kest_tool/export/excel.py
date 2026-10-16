"""Vollständiger Excel-Export (.xlsx) mit allen steuerlich relevanten Daten eines Steuerjahres."""

from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..models import SecurityType
from ..repository import Repository
from ..storage import DocumentStore, sha256_file
from ..tax.engine import (
    CAT_AGE_MANUAL,
    CAT_DIVIDEND,
    CAT_FUND_DIST,
    CAT_FUND_DIST_FREE,
    CAT_INTEREST_25,
    CAT_INTEREST_27,
    CAT_REIT_OTHER,
    CAT_UNCLEAR,
    CAT_CAPREPAY,
    MANUAL,
    PROVISIONAL,
    STATUS_TEXT,
    TaxResult,
)
from ..tax.rules import TaxRules, load_rules

SHEETS = [
    "Steuerübersicht", "Alle Transaktionen", "Dividenden und Ausschüttungen", "Fonds und OeKB", "Nichtmeldefonds",
    "REITs und BDCs", "Verlustausgleich", "Steuerberechnung", "E1kv - FinanzOnline", "Prüfprotokoll",
    "Dokumentenverzeichnis",
]

EUR_FMT = '#,##0.00 "€";[Red]-#,##0.00 "€"'
NUM_FMT = '#,##0.######'
PCT_FMT = "0.00%"
DATE_FMT = "DD.MM.YYYY"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")    # vorläufig / ungeprüft
ERR_FILL = PatternFill("solid", fgColor="F8CBAD")     # manuell prüfen / fehlt
OK_FILL = PatternFill("solid", fgColor="E2EFDA")
TOTAL_FONT = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
MISSING = "FEHLT"
UNCHECKED = "UNGEPRÜFT"


@dataclass
class Col:
    header: str
    key: str | Callable[[dict], Any]
    kind: str = "text"  # text | eur | num | pct | date | int
    width: int = 14
    total: bool = False


def _val(row: dict, key) -> Any:
    return key(row) if callable(key) else row.get(key)


def _to_date(v):
    if isinstance(v, (date, datetime)):
        return v
    if isinstance(v, str) and len(v) >= 10 and v[4] == "-" and v[7] == "-":
        try:
            return date.fromisoformat(v[:10])
        except ValueError:
            return v
    return v


def _status_fill(text: str | None):
    t = (text or "").lower()
    if "manuell" in t or "fehl" in t or "nicht berechnet" in t or "nicht verifiziert" in t or t == "fehler":
        return ERR_FILL
    if any(k in t for k in ("vorläufig", "ungeprüft", "ungeprueft", "sekundär", "prüfen", "pruefen", "teilweise",
                            "neu")) or t == "warnung":
        return WARN_FILL
    if any(k in t for k in ("berechnet", "geprüft", "geprueft", "offiziell", "verarbeitet")):
        return OK_FILL
    return None


class SheetWriter:
    def __init__(self, ws, title: str, subtitle: str = ""):
        self.ws = ws
        ws["A1"] = title
        ws["A1"].font = TITLE_FONT
        if subtitle:
            ws["A2"] = subtitle
            ws["A2"].alignment = Alignment(wrap_text=False)
        self.row = 4
        self.first_table = True

    def note(self, text: str, bold: bool = False) -> None:
        c = self.ws.cell(row=self.row, column=1, value=text)
        if bold:
            c.font = TOTAL_FONT
        self.row += 1

    def table(self, cols: list[Col], rows: list[dict], status_col: str | None = None, totals: bool = True,
              heading: str | None = None, empty_text: str = "Keine Einträge") -> tuple[int, int]:
        ws = self.ws
        if heading:
            ws.cell(row=self.row, column=1, value=heading).font = TOTAL_FONT
            self.row += 1
        hr = self.row
        for i, c in enumerate(cols, 1):
            cell = ws.cell(row=hr, column=i, value=c.header)
            cell.font, cell.fill = HEADER_FONT, HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = Border(bottom=THIN)
            cur = ws.column_dimensions[get_column_letter(i)].width or 0
            ws.column_dimensions[get_column_letter(i)].width = max(cur, c.width)
        r = hr + 1
        if not rows:
            ws.cell(row=r, column=1, value=empty_text).font = Font(italic=True)
            r += 1
        for row in rows:
            for i, c in enumerate(cols, 1):
                v = _val(row, c.key)
                cell = ws.cell(row=r, column=i)
                if c.kind in ("eur", "num", "pct", "int"):
                    if v is None or v == "":
                        cell.value = None
                    else:
                        try:
                            cell.value = float(v) if c.kind != "int" else int(v)
                        except (TypeError, ValueError):
                            cell.value = str(v)
                    cell.number_format = {"eur": EUR_FMT, "num": NUM_FMT, "pct": PCT_FMT, "int": "0"}[c.kind]
                elif c.kind == "date":
                    cell.value = _to_date(v)
                    cell.number_format = DATE_FMT
                else:
                    if isinstance(v, (list, tuple)):
                        v = "; ".join(str(x) for x in v)
                    elif isinstance(v, dict):
                        v = json.dumps(v, ensure_ascii=False, default=str)
                    cell.value = v
                    if isinstance(v, str) and (v == MISSING or v == UNCHECKED):
                        cell.fill = ERR_FILL if v == MISSING else WARN_FILL
                if isinstance(cell.value, str):
                    cell.data_type = "s"
            if status_col:
                idx = [c.header for c in cols].index(status_col) + 1
                fill = _status_fill(ws.cell(row=r, column=idx).value)
                if fill:
                    ws.cell(row=r, column=idx).fill = fill
            r += 1
        last_data = r - 1
        if totals and rows and any(c.total for c in cols):
            ws.cell(row=r, column=1, value="Summe (SUBTOTAL – berücksichtigt Filter)").font = TOTAL_FONT
            for i, c in enumerate(cols, 1):
                if c.total:
                    L = get_column_letter(i)
                    cell = ws.cell(row=r, column=i, value=f"=SUBTOTAL(9,{L}{hr + 1}:{L}{last_data})")
                    cell.font, cell.number_format = TOTAL_FONT, EUR_FMT if c.kind == "eur" else NUM_FMT
                    cell.border = Border(top=THIN)
            r += 1
        if self.first_table:
            ws.freeze_panes = ws.cell(row=hr + 1, column=1)
            ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(cols))}{max(last_data, hr + 1)}"
            self.first_table = False
        self.row = r + 2
        return hr, last_data


# --------------------------------------------------------------------------- #
def _tx_rows(repo: Repository) -> list[dict]:
    rows = repo.transactions(include_deleted=False)
    links = repo.db.query("SELECT transaction_id, group_concat(document_id) docs FROM document_links "
                          "GROUP BY transaction_id")
    linkmap = {l["transaction_id"]: l["docs"] for l in links}
    for r in rows:
        r["Steuerjahr"] = int(r["trade_date"][:4])
        r["Belege"] = linkmap.get(r["id"], "")
        r["EUR-Betrag"] = r["amount_eur"] if r["amount_eur"] is not None else (MISSING if r["amount"] else None)
        r["Quelle"] = f"{r['source_type']}" + (f" / {r['source_ref']}" if r.get("source_ref") else "")
    return rows


def _income_rows(res: TaxResult) -> list[dict]:
    cats = {CAT_DIVIDEND, CAT_FUND_DIST, CAT_FUND_DIST_FREE, CAT_INTEREST_25, CAT_INTEREST_27, CAT_AGE_MANUAL,
            CAT_REIT_OTHER, CAT_UNCLEAR}
    out = []
    for l in res.lines:
        if l.category not in cats:
            continue
        gross = l.gross
        out.append({
            "Datum": l.date, "Depot": l.depot, "Broker": l.broker, "ISIN": l.isin, "Name": l.name,
            "Kategorie": l.category, "Inland/Ausland": "Inland (KESt-Abzug)" if l.domestic else "Ausland",
            "Brutto": gross if gross is not None else MISSING, "Quellensteuer": l.withholding_tax,
            "anrechenbar": l.creditable_tax_max, "nicht anrechenbar": l.non_creditable_tax,
            "ungeklärt": l.unclear_tax, "KESt AT": l.kest_at,
            "Netto": (gross - l.withholding_tax - l.kest_at) if gross is not None else None,
            "Steuerlicher Ansatz": l.amount if l.amount is not None else "nicht berechnet",
            "Steuerjahr": l.tax_year, "Status": STATUS_TEXT[l.status], "Position": l.line_id,
            "Herkunft": "; ".join(l.sources), "Hinweise": "; ".join(l.notes),
        })
    return out


def _calc_rows(res: TaxResult, rules: TaxRules) -> list[dict]:
    out = []
    for l in res.lines:
        out.append({
            "Position": l.line_id, "Steuerjahr": l.tax_year, "Datum": l.date, "Kategorie": l.category,
            "Steuergruppe": l.bucket, "Depot": l.depot, "ISIN": l.isin, "Name": l.name,
            "Ausgangsbetrag": l.gross if l.gross is not None else MISSING,
            "Steuerlicher Ansatz": l.amount if l.amount is not None else "nicht berechnet",
            "Verlustausgleich": -l.offset_amount if l.offset_amount else 0.0,
            "Bemessungsgrundlage": l.taxable_after_offset if l.amount is not None and l.amount > 0 else 0.0,
            "Satz": l.rate, "Steuer (vor Anrechnung)": (l.taxable_after_offset * l.rate) if l.rate and l.amount else 0.0,
            "Regeln": ", ".join(l.rule_ids),
            "Regelstatus": "; ".join(f"{rid}: {rules.rule(rid, l.tax_year).status_label}" for rid in l.rule_ids),
            "Rechenschritte": "\n".join(l.steps), "Quellen": "; ".join(l.sources), "Status": STATUS_TEXT[l.status],
            "Hinweise": "; ".join(l.notes),
        })
    return out


def _loss_rows(res: TaxResult) -> list[dict]:
    out = []
    for r in res.loss_offsets:
        out.append({
            "Verlustposition": r.loss_line, "Steuerjahr": r.tax_year, "Verlustkategorie": r.category,
            "Depot": r.depot, "ISIN": r.isin, "Name": r.name, "Ursprünglicher Verlust": r.original_loss,
            "Zugeordnete Erträge": "; ".join(f"{lid} {cat}: {amt:.2f}" for lid, cat, amt in r.offsets) or "-",
            "Verrechneter Betrag": r.offset_total, "Nicht verrechenbar (verfällt)": r.remaining,
            "Begründungen": "; ".join(r.reasons), "Status": STATUS_TEXT[r.status],
        })
    return out


def _reit_rows(res: TaxResult, repo: Repository) -> list[dict]:
    secs = {s["isin"]: s for s in repo.securities()}
    out = []
    for l in res.lines:
        s = secs.get(l.isin or "", {})
        if s.get("security_type") not in (SecurityType.REIT.value, SecurityType.BDC.value):
            continue
        out.append({
            "Datum": l.date, "Depot": l.depot, "ISIN": l.isin, "Name": l.name, "Art": s.get("security_type"),
            "Ertragsart": l.category, "Bruttoertrag": l.gross, "Quellensteuer": l.withholding_tax,
            "Kapitalrückzahlung": l.gross if l.category == CAT_CAPREPAY else 0.0,
            "Steuerlicher Ansatz": l.amount if l.amount is not None else "nicht berechnet",
            "Steuerliche Klassifizierung": f"{s.get('tax_class')} ({'geprüft' if s.get('tax_class_verified') else UNCHECKED})",
            "Klassifizierungsquelle": s.get("tax_class_source") or MISSING,
            "Prüfstatus": STATUS_TEXT[l.status], "Hinweise": "; ".join(l.notes),
        })
    return out


def _doc_rows(repo: Repository, store: DocumentStore) -> list[dict]:
    rows = repo.documents()
    for d in rows:
        txs = repo.transactions_for_document(d["id"])
        d["Transaktionen"] = ", ".join(str(t["id"]) for t in txs) or "-"
        d["Speicherpfad"] = str(store.base / d["stored_path"])
    return rows


def build_workbook(repo: Repository, store: DocumentStore, res: TaxResult, rules: TaxRules | None = None) -> Workbook:
    rules = rules or load_rules()
    year = res.tax_year
    created = datetime.now().strftime("%d.%m.%Y %H:%M")
    disclaimer = (f"Steuerjahr {year} · erstellt {created} · Arbeitshilfe – keine Steuerberatung, keine "
                  f"offizielle Einreichung · Regelstand {rules.last_review}")
    wb = Workbook()
    wb.remove(wb.active)
    ws = {name: wb.create_sheet(name) for name in SHEETS}

    # 1 Steuerübersicht -----------------------------------------------------
    w = SheetWriter(ws["Steuerübersicht"], f"KESt-Auswertung {year}", disclaimer)
    summ, conf = res.summary, res.summary_confirmed
    hints = {
        "Bemessungsgrundlage 27,5 %": "nach Verlustausgleich",
        "Anrechenbare Quellensteuer (nach Verlustausgleich, Schätzung)": "Schätzung – Bescheid maßgeblich",
        "Geschätzte Steuer nach Anrechnung": "Schätzung; inländische, bereits endbesteuerte Einkünfte enthalten",
        "Ungeklärte Beträge (nicht berechnet)": "nicht in der Bemessungsgrundlage – manuell prüfen",
        "Ungeklärte Quellensteuer": "DBA-Satz/Quellenstaat nicht verifiziert",
    }
    sum_rows = [{"Kennzahl": k, "Gesamt (inkl. vorläufig)": v, "davon gesichert": conf.get(k),
                 "Hinweis": hints.get(k, "")} for k, v in summ.items()]
    w.table([Col("Kennzahl", "Kennzahl", width=58), Col("Gesamt (inkl. vorläufig)", "Gesamt (inkl. vorläufig)",
                                                         "eur", 22),
             Col("davon gesichert", "davon gesichert", "eur", 18), Col("Hinweis", "Hinweis", width=60)],
            sum_rows, totals=False, heading="Ergebnis")
    for r in range(5, w.row):
        label = w.ws.cell(row=r, column=1).value
        if isinstance(label, str) and label.startswith("Positionen"):
            for c in (2, 3):
                w.ws.cell(row=r, column=c).number_format = "0"
    depots = repo.depots()
    dep_rows = []
    for d in depots:
        s = repo.depot_tax_settings(d["id"], year)
        dep_rows.append({"Broker": d["broker_name"], "Depot": d["name"], "Währung": d["currency"],
                         "Österr. KESt-Abzug durch Broker": s["kest_withheld"] if not s.get("_default") else
                         f"{UNCHECKED} (unbekannt)",
                         "Verlustausgleich durch Broker": s["loss_offset_by_broker"],
                         "Quelle der Angabe": s.get("source") or "-"})
    w.table([Col("Broker", "Broker", width=24), Col("Depot", "Depot", width=24), Col("Währung", "Währung", width=9),
             Col("Österr. KESt-Abzug durch Broker", "Österr. KESt-Abzug durch Broker", width=26),
             Col("Verlustausgleich durch Broker", "Verlustausgleich durch Broker", width=24),
             Col("Quelle der Angabe", "Quelle der Angabe", width=30)], dep_rows, totals=False,
            heading="Broker und Depots")
    sev = {}
    for i in res.issues:
        sev[i.severity] = sev.get(i.severity, 0) + 1
    w.table([Col("Schweregrad", "Schweregrad", width=20), Col("Anzahl", "Anzahl", "int", 10)],
            [{"Schweregrad": k, "Anzahl": v} for k, v in sorted(sev.items())], totals=False,
            heading="Offene Prüfungen (Details im Blatt Prüfprotokoll)")
    w.note("Statuslegende: grün = Regel offiziell belegt und Daten geprüft · gelb = vorläufig/ungeprüft/"
           "Sekundärquelle · rot = fehlt bzw. manuell steuerlich zu prüfen (nicht berechnet).")

    # 2 Alle Transaktionen ---------------------------------------------------
    w = SheetWriter(ws["Alle Transaktionen"], "Alle Transaktionen (alle Jahre)", disclaimer)
    w.table([
        Col("ID", "id", "int", 7), Col("Steuerjahr", "Steuerjahr", "int", 10), Col("Broker", "broker_name", width=18),
        Col("Depot", "depot_name", width=18), Col("ISIN", "isin", width=15), Col("Wertpapier", "security_name", width=28),
        Col("Transaktionsart", "tx_type", width=16), Col("Datum", "trade_date", "date", 12),
        Col("Valuta", "value_date", "date", 12), Col("Stückzahl", "quantity", "num", 11), Col("Kurs", "price", "num", 11),
        Col("Währung", "currency", width=8), Col("Betrag (Original)", "amount", "num", 14),
        Col("Gebühren (Original)", "fees", "num", 12), Col("Devisenkurs (je EUR)", "fx_rate", "num", 12),
        Col("Kursquelle", "fx_source", width=28),
        Col("Betrag EUR", "EUR-Betrag", "eur", 14, True), Col("Gebühren EUR", "fees_eur", "eur", 12, True),
        Col("Quellensteuer EUR", "withholding_tax_eur", "eur", 13, True), Col("Quellenstaat", "withholding_country", width=9),
        Col("KESt AT EUR", "kest_at_eur", "eur", 12, True), Col("Ertragsart", "income_category", width=16),
        Col("Prüfstatus", "review_status", width=12), Col("Quelle", "Quelle", width=20),
        Col("Ursprungsdokument", "document_name", width=30), Col("Weitere Belege (Dok.-IDs)", "Belege", width=14),
        Col("Notiz", "notes", width=30),
    ], _tx_rows(repo), status_col="Prüfstatus")

    # 3 Dividenden und Ausschüttungen -------------------------------------------
    w = SheetWriter(ws["Dividenden und Ausschüttungen"], f"Dividenden, Zinsen und Ausschüttungen {year}", disclaimer)
    w.table([
        Col("Datum", "Datum", "date", 12), Col("Depot", "Depot", width=16), Col("Broker", "Broker", width=16),
        Col("ISIN", "ISIN", width=15), Col("Name", "Name", width=26), Col("Steuerliche Kategorie", "Kategorie", width=30),
        Col("Inland/Ausland", "Inland/Ausland", width=16), Col("Bruttoertrag EUR", "Brutto", "eur", 14, True),
        Col("Quellensteuer EUR", "Quellensteuer", "eur", 13, True),
        Col("davon anrechenbar (max.)", "anrechenbar", "eur", 14, True),
        Col("nicht anrechenbar", "nicht anrechenbar", "eur", 13, True), Col("ungeklärt", "ungeklärt", "eur", 12, True),
        Col("Österr. KESt", "KESt AT", "eur", 12, True), Col("Nettoertrag EUR", "Netto", "eur", 14, True),
        Col("Steuerlicher Ansatz", "Steuerlicher Ansatz", "eur", 14, True), Col("Steuerjahr", "Steuerjahr", "int", 9),
        Col("Status", "Status", width=26), Col("Position", "Position", width=14),
        Col("Herkunftsdokument / Quelle", "Herkunft", width=50), Col("Hinweise", "Hinweise", width=50),
    ], _income_rows(res), status_col="Status")

    # 4 Fonds und OeKB ------------------------------------------------------------
    w = SheetWriter(ws["Fonds und OeKB"], f"Fonds und OeKB-Meldungen {year}", disclaimer)
    w.table([
        Col("ISIN", "ISIN", width=15), Col("Fondsname", "Fondsname", width=28), Col("Depot", "Depot", width=16),
        Col("Meldefondsstatus", "Meldefondsstatus", width=14), Col("Meldungsart", "Meldungsart", width=16),
        Col("Geschäftsjahresende", "Geschäftsjahresende", "date", 12), Col("Meldedatum", "Meldedatum", "date", 12),
        Col("Bestand zum Meldedatum", "Bestand zum Meldedatum", "num", 12), Col("Währung", "Währung", width=8),
        Col("Ausschüttung je Anteil", "Ausschüttung je Anteil", "num", 12),
        Col("AgE je Anteil", "AgE je Anteil", "num", 12), Col("AgE gesamt EUR", "AgE gesamt EUR", "eur", 14, True),
        Col("AK-Korrektur je Anteil", "AK-Korrektur je Anteil", "num", 12),
        Col("AK-Korrektur gesamt EUR", "AK-Korrektur gesamt EUR", "eur", 14, True),
        Col("Anrechenbare Steuer je Anteil", "Anrechenbare Steuer je Anteil", "num", 12),
        Col("Quelle", "Quelle", width=40), Col("Version", "Version", "int", 8), Col("Prüfstatus", "Prüfstatus", width=12),
    ], res.fund_rows, status_col="Prüfstatus", heading="Steuerlich angesetzte Fondsmeldungen (Bestand zum Meldedatum)")
    all_reports = repo.fund_reports()
    w.table([
        Col("ID", "id", "int", 6), Col("ISIN", "isin", width=15), Col("Fondsname", "fund_name", width=28),
        Col("Meldungsart", "report_type", width=16), Col("Geschäftsjahresende", "fiscal_year_end", "date", 12),
        Col("Meldedatum", "report_date", "date", 12), Col("Ausschüttung je Anteil", "distribution_per_unit", "num", 12),
        Col("AgE je Anteil", "age_per_unit", "num", 12),
        Col("Steuerpfl. Ausschüttung je Anteil", "taxable_distribution_per_unit", "num", 12),
        Col("AK-Korrektur je Anteil", "ak_correction_per_unit", "num", 12),
        Col("Anrechenbare Steuer je Anteil", "creditable_tax_per_unit", "num", 12), Col("Währung", "currency", width=8),
        Col("Quelle", "source", width=14), Col("Importdatum", "imported_at", width=18), Col("Version", "version", "int", 8),
        Col("Aktiv", lambda r: "ja" if r["active"] else "nein (ersetzt/offen)", width=14),
        Col("Prüfstatus", "review_status", width=12),
    ], all_reports, status_col="Prüfstatus", totals=False, heading="Alle erfassten Fondsmeldungen (alle Versionen)")
    w.note("Automatischer OeKB-Abruf: nicht verfügbar – keine verifizierte, frei nutzbare Schnittstelle. "
           "Werte stammen aus manuellen Importen/Erfassungen.")

    # 5 Nichtmeldefonds ---------------------------------------------------------------
    w = SheetWriter(ws["Nichtmeldefonds"], f"Nichtmeldefonds {year}", disclaimer)
    nmf_cols = ["ISIN", "Name", "Depot", "Steuerjahr", "Bestand 31.12.", "Erster Rücknahmepreis", "Datum erster Preis",
                "Letzter Rücknahmepreis", "Datum letzter Preis", "Währung", "Devisenkurs 31.12.", "Regel", "Regelstatus",
                "Rechenschritte", "Ergebnis EUR", "Datenlücken / Prüfhinweise"]
    kinds = {"Bestand 31.12.": "num", "Erster Rücknahmepreis": "num", "Letzter Rücknahmepreis": "num",
             "Devisenkurs 31.12.": "num", "Ergebnis EUR": "eur", "Steuerjahr": "int",
             "Datum erster Preis": "date", "Datum letzter Preis": "date"}
    w.table([Col(c, c, kinds.get(c, "text"), 40 if c in ("Rechenschritte", "Regel", "Datenlücken / Prüfhinweise")
                 else 14, c == "Ergebnis EUR") for c in nmf_cols], res.nmf_rows, status_col="Regelstatus")

    # 6 REITs und BDCs --------------------------------------------------------------
    w = SheetWriter(ws["REITs und BDCs"], f"REITs und BDCs {year}", disclaimer)
    w.table([
        Col("Datum", "Datum", "date", 12), Col("Depot", "Depot", width=16), Col("ISIN", "ISIN", width=15),
        Col("Name", "Name", width=26), Col("Art", "Art", width=8), Col("Ertragsart", "Ertragsart", width=28),
        Col("Bruttoertrag EUR", "Bruttoertrag", "eur", 14, True), Col("Quellensteuer EUR", "Quellensteuer", "eur", 13, True),
        Col("Kapitalrückzahlung EUR", "Kapitalrückzahlung", "eur", 14, True),
        Col("Steuerlicher Ansatz", "Steuerlicher Ansatz", "eur", 14, True),
        Col("Steuerliche Klassifizierung", "Steuerliche Klassifizierung", width=36),
        Col("Klassifizierungsquelle", "Klassifizierungsquelle", width=30), Col("Prüfstatus", "Prüfstatus", width=26),
        Col("Hinweise", "Hinweise", width=50),
    ], _reit_rows(res, repo), status_col="Prüfstatus")

    # 7 Verlustausgleich --------------------------------------------------------------
    w = SheetWriter(ws["Verlustausgleich"], f"Verlustausgleich {year} (§ 27 Abs 8 EStG)", disclaimer)
    w.table([
        Col("Verlustposition", "Verlustposition", width=16), Col("Steuerjahr", "Steuerjahr", "int", 9),
        Col("Verlustkategorie", "Verlustkategorie", width=20), Col("Depot", "Depot", width=16),
        Col("ISIN", "ISIN", width=15), Col("Name", "Name", width=24),
        Col("Ursprünglicher Verlust", "Ursprünglicher Verlust", "eur", 14, True),
        Col("Zugeordnete Erträge", "Zugeordnete Erträge", width=60),
        Col("Verrechneter Betrag", "Verrechneter Betrag", "eur", 14, True),
        Col("Nicht verrechenbar (verfällt)", "Nicht verrechenbar (verfällt)", "eur", 14, True),
        Col("Begründungen", "Begründungen", width=80), Col("Status", "Status", width=26),
    ], _loss_rows(res), status_col="Status")
    excluded = [l for l in res.lines if l.category == CAT_INTEREST_25 and l.amount]
    if excluded:
        w.note(f"Nicht in den Verlustausgleich einbezogen: {len(excluded)} Zinsposition(en) mit 25 % "
               f"(Summe {sum(l.amount for l in excluded):.2f} EUR) – § 27 Abs 8 Z 1 EStG.")

    # 8 Steuerberechnung --------------------------------------------------------------
    w = SheetWriter(ws["Steuerberechnung"], f"Steuerberechnung {year} – alle Rechenpositionen", disclaimer)
    w.table([
        Col("Position", "Position", width=15), Col("Steuerjahr", "Steuerjahr", "int", 9), Col("Datum", "Datum", "date", 12),
        Col("Steuerkategorie", "Kategorie", width=30), Col("Steuergruppe", "Steuergruppe", width=14),
        Col("Depot", "Depot", width=14), Col("ISIN", "ISIN", width=15), Col("Name", "Name", width=22),
        Col("Ausgangsbetrag EUR", "Ausgangsbetrag", "eur", 14), Col("Steuerlicher Ansatz", "Steuerlicher Ansatz", "eur", 14, True),
        Col("Verlustausgleich", "Verlustausgleich", "eur", 13, True),
        Col("Bemessungsgrundlage", "Bemessungsgrundlage", "eur", 14, True), Col("Steuersatz", "Satz", "pct", 9),
        Col("Steuer vor Anrechnung", "Steuer (vor Anrechnung)", "eur", 14, True), Col("Regeln", "Regeln", width=30),
        Col("Regelstatus", "Regelstatus", width=40), Col("Rechenschritte", "Rechenschritte", width=80),
        Col("Quellen", "Quellen", width=50), Col("Status", "Status", width=26), Col("Hinweise", "Hinweise", width=50),
    ], _calc_rows(res, rules), status_col="Status")
    w.table([Col("Regel-ID", "Regel-ID", width=26), Col("Titel", "Titel", width=50),
             Col("Rechtsgrundlage", "Rechtsgrundlage", width=30), Col("Satz", "Satz", "pct", 8),
             Col("Status", "Status", width=30), Col("Quellen", "Quellen", width=70), Col("Logik", "Logik", width=60),
             Col("Einschränkungen", "Einschränkungen", width=50), Col("Letzte Prüfung", "Letzte Prüfung", width=12)],
            res.rules_used, status_col="Status", totals=False, heading="Angewendete Regeln und Quellen")
    w.table([Col("Bestand", "Depot", width=16), Col("Broker", "Broker", width=16), Col("ISIN", "ISIN", width=15),
             Col("Name", "Name", width=24), Col("Bestand 31.12.", "Bestand 31.12.", "num", 12),
             Col("Steuerliche AK gesamt EUR", "Steuerliche AK gesamt EUR", "eur", 16),
             Col("AK je Stück EUR", "AK je Stück EUR", "num", 12), Col("AK vollständig", "AK vollständig", width=10),
             Col("Kurs 31.12.", "Kurs 31.12.", "num", 12), Col("Kurswährung", "Kurswährung", width=8),
             Col("Kursdatum", "Kursdatum", "date", 12), Col("Kursquelle", "Kursquelle", width=30),
             Col("Wert 31.12. EUR (Info)", "Wert 31.12. EUR", "eur", 16),
             Col("Unrealisiert EUR (Info)", "Unrealisiert EUR (Info)", "eur", 16)],
            res.positions, totals=False, heading=f"Bestände und steuerliche Anschaffungskosten zum 31.12.{year}")

    # 9 E1kv ----------------------------------------------------------------------------
    spec = rules.e1kv(year)
    w = SheetWriter(ws["E1kv - FinanzOnline"], f"E 1kv-Übertragungshilfe {year} – {spec.get('form')}",
                    "Arbeitshilfe zur MANUELLEN Übernahme in FinanzOnline. Keine automatische Übermittlung. "
                    "Beträge vor Übernahme mit Belegen und dem Formular des Steuerjahres abgleichen.")
    w.table([Col("Kennzahl", "Kennzahl", width=10), Col("Bezeichnung", "Bezeichnung", width=60),
             Col("Betrag EUR", "Betrag EUR", "eur", 14), Col("Kennzahl-Status", "Kennzahl-Status", width=24),
             Col("Betrags-Status", "Betrags-Status", width=36), Col("Positionen", "Positionen", width=40),
             Col("Hinweise", "Hinweise", width=90)], res.e1kv, status_col="Betrags-Status", totals=False)
    for n in spec.get("notes", []):
        w.note("Hinweis: " + n)
    missing = [i for i in res.issues if i.kind == "Fehlende Daten"]
    manual = [l for l in res.lines if l.status == MANUAL]
    w.note(f"Fehlende Angaben: {len(missing)} · Manuell zu prüfende Positionen: {len(manual)} "
           "(nicht in den Beträgen enthalten, siehe Prüfprotokoll/Steuerberechnung)", bold=True)
    w.table([Col("Position", "Position", width=16), Col("Kategorie", "Kategorie", width=30),
             Col("Ausgangsbetrag", "Betrag", "eur", 14), Col("Grund", "Grund", width=90)],
            [{"Position": l.line_id, "Kategorie": l.category, "Betrag": l.gross, "Grund": "; ".join(l.notes)}
             for l in manual], totals=False, heading="Manuelle Prüfungen (nicht in E 1kv-Beträgen enthalten)")

    # 10 Prüfprotokoll -----------------------------------------------------------------
    w = SheetWriter(ws["Prüfprotokoll"], f"Prüfprotokoll {year}", disclaimer)
    w.table([Col("Schweregrad", "severity", width=12), Col("Art", "kind", width=24), Col("Meldung", "message", width=100),
             Col("Was tun?", "action", width=70),
             Col("Bezug", "ref", width=50), Col("Steuerjahr", "tax_year", "int", 9)],
            [i.__dict__ for i in res.issues], status_col="Schweregrad", totals=False)
    w.table([Col("ID", "id", "int", 6), Col("Objekt", "entity", width=14), Col("Schlüssel", "entity_key", width=26),
             Col("Feld", "field", width=20), Col("Bestehend", "existing_value", width=16),
             Col("Quelle bestehend", "existing_ref", width=26), Col("Neu", "new_value", width=16),
             Col("Quelle neu", "new_ref", width=26), Col("Status", "status", width=16),
             Col("Begründung", "decision_note", width=40)],
            repo.conflicts(), totals=False, heading="Widersprüchliche Daten (Konflikte)")
    w.table([Col("Zeit", "ts", width=19), Col("Objekt", "entity", width=16), Col("ID", "entity_id", width=10),
             Col("Aktion", "action", width=20), Col("Details", "details", width=120)],
            repo.audit_entries(), totals=False, heading="Manuelle Änderungen und Protokoll (Audit-Log)")

    # 11 Dokumentenverzeichnis ------------------------------------------------------
    w = SheetWriter(ws["Dokumentenverzeichnis"], "Dokumentenverzeichnis (alle Dokumente)", disclaimer)
    w.table([Col("ID", "id", "int", 6), Col("Dateiname", "original_filename", width=40), Col("Dateityp", "file_type", width=8),
             Col("Dokumenttyp", "doc_type", width=22), Col("Broker/Quelle", "source", width=16),
             Col("Steuerjahr", "tax_year", "int", 9), Col("Importdatum", "imported_at", width=19),
             Col("SHA-256", "sha256", width=66), Col("Zugehörige Transaktionen", "Transaktionen", width=20),
             Col("Speicherpfad", "Speicherpfad", width=80), Col("Verarbeitungsstatus", "status", width=14)],
            _doc_rows(repo, store), status_col="Verarbeitungsstatus", totals=False)
    for sheet in wb.worksheets:
        sheet.sheet_view.zoomScale = 90
    return wb


def export_excel(repo: Repository, store: DocumentStore, res: TaxResult, rules: TaxRules | None = None,
                 today: date | None = None) -> Path:
    """Erzeugt die Datei ``KESt_Auswertung_<Jahr>_<Exportdatum>.xlsx`` im Ordner exports/<Jahr>/excel."""
    today = today or date.today()
    store.ensure_year(res.tax_year)
    wb = build_workbook(repo, store, res, rules)
    target_dir = store.export_dir(res.tax_year, "excel")
    target = store.unique_target(target_dir, f"KESt_Auswertung_{res.tax_year}_{today.isoformat()}.xlsx")
    tmp = target.with_suffix(".xlsx.part")
    wb.save(tmp)
    tmp.replace(target)
    repo.record_export(res.tax_year, "excel", target, sha256_file(target),
                       f"{len(res.lines)} Positionen, {len(res.issues)} Prüfhinweise")
    return target
